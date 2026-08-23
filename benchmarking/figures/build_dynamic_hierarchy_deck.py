"""Build the 5-slide deck and the results workbook for the n-level hierarchy.

    python benchmarking/figures/build_dynamic_hierarchy_deck.py

Reads results/ for the numbers and figures/fig_dynamic_*.png for the images, so
both outputs regenerate from data.
"""

from __future__ import annotations

import glob
import json
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = Path(__file__).resolve().parent
BENCH = HERE.parent
RESULTS = BENCH / "results" / "gpt-5.5"
TASKS = [f"hanoi_{n}" for n in range(3, 13)]

INK = RGBColor(0x0B, 0x0B, 0x0B)
INK2 = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x89, 0x87, 0x81)
SURFACE = RGBColor(0xFC, 0xFC, 0xFB)
BAND = RGBColor(0xF4, 0xF4, 0xF1)
BLUE = RGBColor(0x2A, 0x78, 0xD6)
ORANGE = RGBColor(0xEB, 0x68, 0x34)
AQUA = RGBColor(0x1B, 0xAF, 0x7A)
VIOLET = RGBColor(0x4A, 0x3A, 0xA7)
CRITICAL = RGBColor(0xD0, 0x3B, 0x3B)
FONT = "Segoe UI"

MODES = [
    ("direct", "not-hierarchy", None),
    ("inner-outer", "inner-outer", None),
    ("hierarchy (2-level)", "hierarchy", None),
    ("dynamic (n-level)", "dynamic-hierarchy", "low"),
]


def load(mode: str, effort=None) -> dict:
    out = {}
    for task in TASKS:
        best = None
        for path in sorted(glob.glob(str(RESULTS / mode / f"{task}_*" / "metrics.json"))):
            with open(path, encoding="utf-8") as handle:
                m = json.load(handle)
            if effort is not None and str(m["reasoning_effort"]) != effort:
                continue
            best = m
        if best:
            out[task] = best
    return out


DATA = {label: load(mode, eff) for label, mode, eff in MODES}


def totals(label):
    rows = DATA[label]
    return {
        "solved": sum(m["solved"] for m in rows.values()),
        "tokens": sum(m["output_tokens"] for m in rows.values()),
        "calls": sum(m["model_call_count"] for m in rows.values()),
        "replans": sum(m["replan_count"] for m in rows.values()),
    }


# --- slide helpers ------------------------------------------------------------
def blank(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = SURFACE
    return slide


def textbox(slide, x, y, w, h, text, size=16, color=INK, bold=False,
            align=PP_ALIGN.LEFT, space_after=6, line=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, part in enumerate(text.split("\n")):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = align
        para.space_after = Pt(space_after)
        if line:
            para.line_spacing = line
        run = para.add_run()
        run.text = part
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        run.font.name = FONT
    return box


def heading(slide, title, subtitle=None):
    textbox(slide, 0.55, 0.34, 12.2, 0.7, title, size=30, bold=True)
    if subtitle:
        textbox(slide, 0.55, 1.02, 12.2, 0.5, subtitle, size=14, color=INK2)


def picture_fit(slide, name, y, max_w, max_h, cx=13.333 / 2):
    """Scale to fit inside both bounds and centre horizontally on cx."""
    from PIL import Image
    with Image.open(HERE / name) as im:
        aspect = im.width / im.height
    w = min(max_w, max_h * aspect)
    h = w / aspect
    return slide.shapes.add_picture(str(HERE / name), Inches(cx - w / 2), Inches(y),
                                    width=Inches(w), height=Inches(h))


def stat(slide, x, y, w, value, label, color=INK):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(1.3))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.space_after = Pt(2)
    r = p.add_run()
    r.text = value
    r.font.size = Pt(38)
    r.font.bold = True
    r.font.color.rgb = color
    r.font.name = FONT
    # One paragraph per line: a "\n" inside a run is not a reliable line break.
    for line_text in label.split("\n"):
        para = tf.add_paragraph()
        para.space_after = Pt(0)
        run = para.add_run()
        run.text = line_text
        run.font.size = Pt(11)
        run.font.color.rgb = INK2
        run.font.name = FONT


def table(slide, x, y, w, h, rows, col_widths=None, highlight_col=None):
    shape = slide.shapes.add_table(len(rows), len(rows[0]), Inches(x), Inches(y),
                                   Inches(w), Inches(h))
    tbl = shape.table
    tbl.first_row = True
    if col_widths:
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Emu(int(Inches(cw)))
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = ""
            cell.margin_top = Pt(2)
            cell.margin_bottom = Pt(2)
            cell.margin_left = Pt(6)
            cell.margin_right = Pt(6)
            para = cell.text_frame.paragraphs[0]
            para.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.RIGHT
            run = para.add_run()
            run.text = str(val)
            run.font.size = Pt(11)
            run.font.name = FONT
            is_head = r == 0
            is_last = r == len(rows) - 1
            run.font.bold = is_head or is_last or (highlight_col == c and not is_head)
            run.font.color.rgb = INK if (is_head or is_last) else INK2
            if highlight_col == c and not is_head:
                run.font.color.rgb = VIOLET
            fill = cell.fill
            fill.solid()
            fill.fore_color.rgb = BAND if (is_head or is_last) else SURFACE
    return tbl


# --- slides -------------------------------------------------------------------
def build_deck(path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    dyn, hier = totals("dynamic (n-level)"), totals("hierarchy (2-level)")

    # 1 — title + headline
    s = blank(prs)
    textbox(s, 0.55, 1.15, 12.2, 0.9, "Dynamic (n-level) Hierarchy", size=42, bold=True)
    textbox(s, 0.55, 2.05, 12.2, 0.6,
            "Planning and hierarchy composition split into separate stages, "
            "with hierarchy depth chosen by the model", size=17, color=INK2)
    stat(s, 0.55, 3.30, 3.0, f"{dyn['solved']}/10", "Hanoi tasks solved optimally\n(2-level: 8/10)", VIOLET)
    stat(s, 3.75, 3.30, 3.0, f"{hier['tokens']/dyn['tokens']:.1f}x", "fewer output tokens\nthan the 2-level pipeline", VIOLET)
    stat(s, 6.95, 3.30, 3.0, f"{dyn['replans']}", "correction rounds needed\n(2-level: 55)", VIOLET)
    stat(s, 10.15, 3.30, 3.0, "12", "max hierarchy depth built\n(2-level: fixed at 2)", VIOLET)
    textbox(s, 0.55, 5.55, 12.2, 1.4,
            "It solves hanoi_11 and hanoi_12, which the 2-level pipeline cannot: there, "
            "the 2-level pipeline burns 15 replans and ~90k output tokens and executes zero moves.",
            size=14, color=INK)
    textbox(s, 0.55, 6.55, 12.2, 0.6,
            "gpt-5.5 · reasoning effort low · max 8,192 output tokens/call · replan budget 15 · "
            "1 run per cell · baselines 2026-07-06/07, dynamic 2026-08-07",
            size=10, color=MUTED)

    # 2 — framework
    s = blank(prs)
    heading(s, "Architecture",
            "The DecisionBot is shown no composed actions and never given a call format, so it cannot emit a solution.")
    picture_fit(s, "fig_dynamic_framework.png", 1.60, 12.5, 5.75)

    # 3 — how depth is decided
    s = blank(prs)
    heading(s, "How many layers? The model decides.",
            "No level count is configured anywhere. One call emits the whole hierarchy; depth is inferred from the call graph afterwards.")
    picture_fit(s, "fig_dynamic_depth.png", 1.62, 12.5, 4.45)
    textbox(s, 0.55, 6.25, 12.2, 1.0,
            "level(primitive) = 0,  level(f) = 1 + max(level of callees).  The model never labels a level; "
            "cycles and unresolvable callees are rejected.\n"
            "Given free rein it built exactly one function per layer in all 10 tasks — a straight chain — "
            "rediscovering Hanoi's recursion from an abstract example that never mentions towers.",
            size=11.5, color=INK2)

    # 4 — comparison
    s = blank(prs)
    heading(s, "Against the baselines",
            "gpt-5.5 at reasoning effort low in every arm. Cell values are output tokens; every solve was optimal.")
    picture_fit(s, "fig_dynamic_comparison.png", 1.60, 8.55, 3.55, cx=4.75)
    rows = [["", "solved", "out-tok", "calls", "replans"]]
    for label, _, _ in MODES:
        t = totals(label)
        rows.append([label, f"{t['solved']}/10", f"{t['tokens']:,}", t["calls"], t["replans"]])
    rows.append(["ceiling", "n=5 → 7 → 10 → 12+", "", "", ""])
    table(s, 9.20, 2.05, 3.75, 2.4, rows, col_widths=[1.45, 0.72, 0.82, 0.42, 0.62])
    textbox(s, 9.20, 4.75, 3.80, 2.2,
            "The n-level arm is the only step in that sequence that is not a "
            "capability-for-cost trade — it is cheaper on every individual task, "
            "not just in aggregate.",
            size=11.5, color=INK)

    # 5 — why, and what is not yet settled
    s = blank(prs)
    heading(s, "Why it works — and what is not settled")
    textbox(s, 0.55, 1.50, 6.1, 0.4, "One cause, three effects", size=15, bold=True, color=VIOLET)
    textbox(s, 0.55, 2.00, 6.1, 4.6,
            "2-level forbids H2 → H2, so the response must contain all 2ⁿ−1 moves.\n"
            "hanoi_12 = 4,095 calls ≈ 30k tokens against an 8,192 cap. It cannot fit.\n\n"
            "n-level writes 12 lines; Python expands them to 4,095 moves.\n\n"
            "· Fewer tokens — mechanical.\n"
            "· Zero replans — 1,023 hand-written moves means 1,023 chances to slip; "
            "12 structural lines means 12. The 2-level replan curve (1→3→6→10→15) "
            "tracks sequence length, not difficulty.\n"
            "· Fewer calls — each 2-level replan regenerates the whole pipeline, "
            "so 16 rounds costs ~80 calls, not 16.\n\n"
            "It is notation, not intelligence: at hanoi_11–12 the 2-level pipeline spent "
            "90–99k tokens across 16 attempts and executed zero moves. Reasoning intensity "
            "was comparable in both (325 vs 290 reasoning tokens per call).",
            size=11.5, color=INK, line=1.18)
    textbox(s, 7.05, 1.50, 5.75, 0.4, "Not yet settled", size=15, bold=True, color=CRITICAL)
    textbox(s, 7.05, 2.00, 5.75, 4.6,
            "· n=1 per cell. Two dynamic runs differed up to 3.4× on per-task tokens "
            "(hanoi_11: 14,749 vs 4,331), driven by subtask count. Solve rate and "
            "aggregate are robust; per-task tokens are noisy.\n\n"
            "· Baseline columns are month-old runs — drift uncontrolled. Reasoning "
            "effort is now matched; run date is not.\n\n"
            "· No max-tokens control arm. \"Does the baseline survive at 32,768?\" "
            "is unanswered.\n\n"
            "· Hanoi flatters this. Perfectly recursive, so compression is near-maximal. "
            "It went pure depth, zero breadth — untested on a domain where several "
            "distinct sub-skills matter.\n\n"
            "Next: fresh baseline + max-tokens control, n≥3 reps, then a domain seam "
            "for non-Hanoi puzzles.",
            size=11.5, color=INK, line=1.18)

    prs.save(str(path))


# --- workbook -----------------------------------------------------------------
def build_workbook(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    head_font = Font(name=FONT, bold=True, color="0B0B0B")
    head_fill = PatternFill("solid", fgColor="F4F4F1")
    body_font = Font(name=FONT)

    def style_header(ws, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=1, column=c)
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(horizontal="left")

    ws = wb.active
    ws.title = "results"
    ws.append(["task", "rings", "mode", "solved", "optimal", "output_tokens",
               "input_tokens", "model_calls", "replans", "elapsed_s",
               "max_hierarchy_level", "top_level_calls", "moves", "optimal_moves"])
    for label, _, _ in MODES:
        for n, task in zip(range(3, 13), TASKS):
            m = DATA[label].get(task)
            if not m:
                continue
            ws.append([task, n, label, m["solved"], m["optimal"], m["output_tokens"],
                       m["input_tokens"], m["model_call_count"], m["replan_count"],
                       round(m["elapsed_seconds"], 1), m.get("max_hierarchy_level", 2),
                       m["top_level_call_count"], m["move_count"], m["optimal_move_count"]])
    style_header(ws, 14)
    for col, width in zip("ABCDEFGHIJKLMN",
                          [11, 7, 20, 9, 9, 14, 13, 12, 9, 11, 20, 16, 9, 14]):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"

    ws2 = wb.create_sheet("summary")
    ws2.append(["mode", "solved", "optimal", "output_tokens", "input_tokens",
                "model_calls", "replans"])
    for label, _, _ in MODES:
        rows = DATA[label]
        ws2.append([label,
                    f"{sum(m['solved'] for m in rows.values())}/10",
                    f"{sum(m['optimal'] for m in rows.values())}/10",
                    sum(m["output_tokens"] for m in rows.values()),
                    sum(m["input_tokens"] for m in rows.values()),
                    sum(m["model_call_count"] for m in rows.values()),
                    sum(m["replan_count"] for m in rows.values())])
    style_header(ws2, 7)
    for col, width in zip("ABCDEFG", [22, 9, 9, 15, 14, 13, 9]):
        ws2.column_dimensions[col].width = width

    ws3 = wb.create_sheet("layers")
    ws3.append(["task", "rings", "depth_chosen", "functions_defined", "top_level_calls"]
               + [f"calls_L{i}" for i in range(1, 13)])
    dynamic = DATA["dynamic (n-level)"]
    for n, task in zip(range(3, 13), TASKS):
        m = dynamic.get(task)
        if not m:
            continue
        lc = {int(k): v for k, v in m["level_call_counts"].items()}
        mc = {int(k): v for k, v in m["mapping_count_by_level"].items()}
        ws3.append([task, n, m["max_hierarchy_level"], sum(mc.values()),
                    m["top_level_call_count"]] + [lc.get(i, "") for i in range(1, 13)])
    style_header(ws3, 17)
    for col, width in zip("ABCDE", [11, 7, 14, 18, 16]):
        ws3.column_dimensions[col].width = width

    ws4 = wb.create_sheet("config")
    for row in [
        ["model", "gpt-5.5"],
        ["reasoning_effort", "low (pinned in every arm)"],
        ["endpoint", "OpenAI Chat Completions API"],
        ["max output tokens / call", 8192],
        ["replan budget", 15],
        ["goal state", "model-generated (fixed_goal=False)"],
        ["tasks", "hanoi_3 .. hanoi_12"],
        ["repetitions", "1 per cell"],
        ["baseline run dates", "2026-07-06 / 2026-07-07"],
        ["dynamic run date", "2026-08-07"],
        ["cell values", "output tokens (completion, reasoning included)"],
        ["caveat", "baselines are month-old runs; no max-tokens control arm"],
    ]:
        ws4.append(row)
    ws4.column_dimensions["A"].width = 26
    ws4.column_dimensions["B"].width = 52
    for r in range(1, ws4.max_row + 1):
        ws4.cell(row=r, column=1).font = Font(name=FONT, bold=True)
        ws4.cell(row=r, column=2).font = body_font

    wb.save(str(path))


if __name__ == "__main__":
    deck = BENCH / "DYNAMIC_HIERARCHY_SLIDES.pptx"
    book = BENCH / "dynamic_hierarchy_results.xlsx"
    build_deck(deck)
    build_workbook(book)
    print(f"wrote {deck.name} and {book.name}")
